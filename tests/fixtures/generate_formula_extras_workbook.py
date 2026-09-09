"""Generates a second spreadsheet fixture, sibling to generate_test_workbook.py,
specifically to exercise the formula patterns added to
adapters/spreadsheet/formula.py beyond the original SUM/cross-sheet-cell/one-
"+"-combination set: AVERAGE, COUNT, MAX, MIN, a multi-range "SUM+SUM", and a
minimal IF(condition, true, false).

Produces tests/fixtures/formula_extras_test.xlsx: one sheet, a merged title,
a header row, three literal data cells (B3:B5 = 10, 20, 30), and one formula
cell per new pattern, each referencing that same B3:B5 range so their
expected values are easy to hand-verify:
  AVERAGE(B3:B5) = 20
  COUNT(B3:B5)   = 3
  MAX(B3:B5)     = 30
  MIN(B3:B5)     = 10
  SUM(B3:B4)+SUM(B5:B5) = 30 + 30 = 60   (multi-range "+")
  IF(B3>5, B4, B5) = B4 = 20             (B3=10 > 5, so the true branch)
"""

from pathlib import Path

from openpyxl import Workbook

OUTPUT_PATH = Path(__file__).parent / "formula_extras_test.xlsx"


def build_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "Extras"

    ws["A1"] = "Formula Extras"
    ws.merge_cells("A1:B1")

    ws["A2"] = "Label"
    ws["B2"] = "Value"

    ws["A3"] = "Data1"
    ws["B3"] = 10
    ws["A4"] = "Data2"
    ws["B4"] = 20
    ws["A5"] = "Data3"
    ws["B5"] = 30

    ws["A6"] = "Average"
    ws["B6"] = "=AVERAGE(B3:B5)"
    ws["A7"] = "Count"
    ws["B7"] = "=COUNT(B3:B5)"
    ws["A8"] = "Max"
    ws["B8"] = "=MAX(B3:B5)"
    ws["A9"] = "Min"
    ws["B9"] = "=MIN(B3:B5)"
    ws["A10"] = "Multi-range sum"
    ws["B10"] = "=SUM(B3:B4)+SUM(B5:B5)"
    ws["A11"] = "Conditional"
    ws["B11"] = "=IF(B3>5,B4,B5)"

    wb.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    path = build_workbook()
    print(f"Wrote {path}")
