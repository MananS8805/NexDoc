"""Dedicated unit tests for the spreadsheet adapter — confidence-tier
correctness per formula type, both real checks' pass/fail behavior, and one
test per formula pattern added beyond the original SUM/cross-sheet-cell/
single-"+"-combination set (see adapters/spreadsheet/formula.py's
docstring). Until now the spreadsheet adapter (Phase 1) had zero
adapter-specific tests of its own — only indirect coverage via
tests/test_schema_conformance.py and the eval scripts — unlike the diagram
adapter's three dedicated test files.
"""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

BUDGET_FIXTURE = ROOT / "tests" / "fixtures" / "budget_test.xlsx"
EXTRAS_FIXTURE = ROOT / "tests" / "fixtures" / "formula_extras_test.xlsx"


def _build(path):
    from adapters.spreadsheet.adapter import build_artifact

    return build_artifact(path)


def _regions_by_id(artifact):
    return {r["id"]: r for r in artifact["regions"]}


def test_confidence_tiers_on_budget_fixture():
    artifact = _build(BUDGET_FIXTURE)
    regions = _regions_by_id(artifact)

    # Literal cells: 0.99, a direct unambiguous read.
    assert regions["Actuals!A1"]["confidence"] == 0.99  # title
    assert regions["Actuals!A2"]["confidence"] == 0.99  # header
    assert regions["Actuals!B3"]["confidence"] == 0.99  # value

    # Same-sheet formula: 0.96.
    assert regions["Actuals!E3"]["confidence"] == 0.96
    assert regions["Actuals!E3"]["text"] == "=SUM(B3:D3) → 3780"

    # Cross-sheet formula that DOES resolve: 0.80.
    assert regions["Q1 Plan!E6"]["confidence"] == 0.80

    # Formula that can't be resolved at all (target sheet doesn't exist): 0.40.
    assert regions["Q1 Plan!F3"]["confidence"] == 0.40
    assert "unresolved" in regions["Q1 Plan!F3"]["text"]


def test_cross_sheet_check_fails_on_the_deliberately_broken_reference():
    artifact = _build(BUDGET_FIXTURE)
    check = next(c for c in artifact["checks"] if c["name"] == "cross-sheet-references-resolve")
    assert check["passed"] is False
    assert "Forecast Q4" in check["detail"]


def test_totals_check_passes_on_real_sum_formulas():
    artifact = _build(BUDGET_FIXTURE)
    check = next(c for c in artifact["checks"] if c["name"] == "totals-match-referenced-cells")
    assert check["passed"] is True
    assert "12 totals formula(s) checked" in check["detail"]


def test_average_count_max_min_resolve_at_same_sheet_confidence():
    artifact = _build(EXTRAS_FIXTURE)
    regions = _regions_by_id(artifact)

    assert regions["Extras!B6"]["text"] == "=AVERAGE(B3:B5) → 20.0"
    assert regions["Extras!B7"]["text"] == "=COUNT(B3:B5) → 3"
    assert regions["Extras!B8"]["text"] == "=MAX(B3:B5) → 30"
    assert regions["Extras!B9"]["text"] == "=MIN(B3:B5) → 10"
    for ref in ("B6", "B7", "B8", "B9"):
        assert regions[f"Extras!{ref}"]["confidence"] == 0.96, f"{ref} should resolve at same-sheet confidence"


def test_multi_range_sum_plus_sum_resolves():
    artifact = _build(EXTRAS_FIXTURE)
    region = _regions_by_id(artifact)["Extras!B10"]
    assert region["text"] == "=SUM(B3:B4)+SUM(B5:B5) → 60"
    assert region["confidence"] == 0.96


def test_minimal_if_formula_resolves_the_true_branch():
    artifact = _build(EXTRAS_FIXTURE)
    region = _regions_by_id(artifact)["Extras!B11"]
    # B3=10 > 5, so IF(B3>5, B4, B5) should take the true branch, B4=20.
    assert region["text"] == "=IF(B3>5,B4,B5) → 20"
    assert region["confidence"] == 0.96


def test_unsupported_formula_still_degrades_honestly_not_a_guess():
    """A formula pattern this evaluator was never taught (nested IF,
    string functions, arithmetic beyond a single '+') must still resolve to
    "unresolved (unsupported formula pattern)" at 0.40 confidence — never a
    wrong guessed value. Verified directly rather than assumed.
    """
    import openpyxl

    from adapters.spreadsheet.formula import UNRESOLVED_UNSUPPORTED_PATTERN, evaluate

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = 5
    ws["A2"] = '=CONCATENATE(A1,"x")'  # a function this evaluator has never supported

    resolved = evaluate(wb, ws.title, "A2")
    assert resolved.ok is False
    assert resolved.reason == UNRESOLVED_UNSUPPORTED_PATTERN


if __name__ == "__main__":
    test_confidence_tiers_on_budget_fixture()
    test_cross_sheet_check_fails_on_the_deliberately_broken_reference()
    test_totals_check_passes_on_real_sum_formulas()
    test_average_count_max_min_resolve_at_same_sheet_confidence()
    test_multi_range_sum_plus_sum_resolves()
    test_minimal_if_formula_resolves_the_true_branch()
    test_unsupported_formula_still_degrades_honestly_not_a_guess()
    print("all spreadsheet adapter tests passed")
